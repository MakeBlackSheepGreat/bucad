from __future__ import annotations

import json
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable


CORE_MARKDOWN_REPORTS = [
    "final_validation.md",
    "report_tables.md",
    "threshold_analysis.md",
    "visual_evidence_review.md",
    "release_v1_manifest.md",
    "model_freeze_decision.md",
    "comparison_summary.md",
    "efficientnetv2_s_5fold_summary.md",
]

CORE_JSON_REPORTS = [
    "comparison_results.json",
    "busi_eval.json",
    "busi_eval_final.json",
    "train_cls_fold1.json",
    "train_seg_fold1.json",
    "busbra_split_summary.json",
]

SUPPORTED_FORMATS = {"docx"}


@dataclass
class ReportBlock:
    kind: str
    text: str = ""
    level: int = 1
    rows: list[list[str]] = field(default_factory=list)


def _clean_markdown_inline(text: str) -> str:
    return text.replace("`", "")


def _split_table_row(line: str) -> list[str]:
    return [_clean_markdown_inline(cell.strip()) for cell in line.strip().strip("|").split("|")]


def _is_table_separator(line: str) -> bool:
    stripped = line.replace("|", "").replace(":", "").replace("-", "").strip()
    return not stripped


def _markdown_to_blocks(text: str) -> list[ReportBlock]:
    blocks: list[ReportBlock] = []
    lines = text.splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index].rstrip()
        if line.startswith("```"):
            if in_code:
                if code_lines:
                    blocks.append(ReportBlock("code", "\n".join(code_lines)))
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        if line.startswith("|") and "|" in line[1:]:
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                current = lines[index].rstrip()
                if not _is_table_separator(current):
                    rows.append(_split_table_row(current))
                index += 1
            if rows:
                blocks.append(ReportBlock("table", rows=rows))
            continue
        if line.startswith("#"):
            marker, _, title = line.partition(" ")
            blocks.append(ReportBlock("heading", _clean_markdown_inline(title.strip()), level=min(len(marker), 3)))
        elif line.startswith("- [") or line.startswith("- "):
            blocks.append(ReportBlock("bullet", _clean_markdown_inline(line[2:].strip())))
        else:
            blocks.append(ReportBlock("paragraph", _clean_markdown_inline(line.strip())))
        index += 1
    return blocks


def _format_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.4f}"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _metrics_table(metrics: dict[str, Any]) -> ReportBlock:
    rows = [["指标", "数值"]]
    preferred_keys = (
        "auc",
        "threshold",
        "accuracy",
        "sensitivity",
        "recall",
        "precision",
        "specificity",
        "f1_score",
        "youden_j",
    )
    for key in preferred_keys:
        if key in metrics:
            rows.append([key, _format_value(metrics[key])])
    for key, value in metrics.items():
        if key not in preferred_keys and key != "confusion":
            rows.append([key, _format_value(value)])
    if "confusion" in metrics and isinstance(metrics["confusion"], dict):
        for key, value in metrics["confusion"].items():
            rows.append([f"confusion.{key}", _format_value(value)])
    return ReportBlock("table", rows=rows)


def _comparison_table(results: list[Any]) -> ReportBlock:
    rows = [["模型", "状态", "AUC", "Accuracy", "Sensitivity", "Precision", "Specificity", "F1-Score", "运行时间(秒)"]]
    for item in results[:20]:
        metrics = item.get("metrics", {}) if isinstance(item, dict) else {}
        rows.append(
            [
                _format_value(item.get("model_name", "")),
                _format_value(item.get("status", "")),
                _format_value(metrics.get("auc", "")),
                _format_value(metrics.get("accuracy", "")),
                _format_value(metrics.get("sensitivity", "")),
                _format_value(metrics.get("precision", "")),
                _format_value(metrics.get("specificity", "")),
                _format_value(metrics.get("f1_score", "")),
                _format_value(item.get("runtime_seconds", "")),
            ]
        )
    return ReportBlock("table", rows=rows)


def _json_to_blocks(path: Path, payload: dict[str, Any]) -> list[ReportBlock]:
    blocks = [ReportBlock("heading", f"结构化报告：{path.name}", level=2)]
    if "sample_count" in payload:
        blocks.append(ReportBlock("paragraph", f"样本数量：{payload['sample_count']}"))
    if "fold" in payload:
        blocks.append(ReportBlock("paragraph", f"训练/验证折：{payload['fold']}"))
    summary_rows = [["字段", "值"]]
    for key in (
        "fold_count",
        "unique_cases",
        "leakage_detected",
        "train_size",
        "val_size",
        "device",
        "checkpoint_path",
        "output_path",
        "dry_run",
        "model_count",
    ):
        if key in payload:
            summary_rows.append([key, _format_value(payload[key])])
    if len(summary_rows) > 1:
        blocks.append(ReportBlock("heading", "摘要字段", level=3))
        blocks.append(ReportBlock("table", rows=summary_rows))
    if isinstance(payload.get("metrics"), dict):
        blocks.append(ReportBlock("heading", "核心指标", level=3))
        blocks.append(_metrics_table(payload["metrics"]))
    threshold_analysis = payload.get("threshold_analysis", {})
    if isinstance(threshold_analysis, dict) and isinstance(threshold_analysis.get("best_by_youden"), dict):
        blocks.append(ReportBlock("heading", "Youden J 最优阈值", level=3))
        blocks.append(_metrics_table(threshold_analysis["best_by_youden"]))
    results = payload.get("results")
    if isinstance(results, list):
        blocks.append(ReportBlock("heading", "模型对比结果", level=3))
        blocks.append(_comparison_table(results))
        if len(results) > 20:
            blocks.append(ReportBlock("paragraph", f"为保证版面可读性，仅展示前 20 条，共 {len(results)} 条。"))
    elif not any(key in payload for key in ("metrics", "threshold_analysis", "sample_count", "fold")):
        rows = [["字段", "值"]]
        for key, value in payload.items():
            rows.append([str(key), _format_value(value)])
        blocks.append(ReportBlock("table", rows=rows))
    return blocks


def collect_report_blocks(reports_dir: str | Path) -> list[ReportBlock]:
    root = Path(reports_dir)
    if not root.exists():
        raise FileNotFoundError(f"Report directory does not exist: {root}")

    source_files: list[Path] = []
    for name in CORE_MARKDOWN_REPORTS + CORE_JSON_REPORTS:
        path = root / name
        if path.exists():
            source_files.append(path)

    included = ", ".join(path.name for path in source_files) if source_files else "未找到可汇总的核心报告文件"
    blocks = [
        ReportBlock("title", "BUCAD 项目报告汇总"),
        ReportBlock("paragraph", f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"),
        ReportBlock("paragraph", f"报告目录：{root.resolve()}"),
        ReportBlock("heading", "说明", level=1),
        ReportBlock(
            "paragraph",
            "本文件汇总项目主要验证、评估、可视化和发布相关报告，便于阅读和归档。",
        ),
        ReportBlock("bullet", f"纳入文件：{included}"),
    ]

    for name in CORE_MARKDOWN_REPORTS:
        path = root / name
        if path.exists():
            blocks.append(ReportBlock("heading", f"源报告：{path.name}", level=1))
            blocks.extend(_markdown_to_blocks(path.read_text(encoding="utf-8")))
    for name in CORE_JSON_REPORTS:
        path = root / name
        if path.exists():
            with path.open("r", encoding="utf-8") as handle:
                payload = json.load(handle)
            if isinstance(payload, dict):
                blocks.extend(_json_to_blocks(path, payload))
    return blocks


def _iter_table_rows(rows: list[list[str]]) -> Iterable[list[str]]:
    width = max((len(row) for row in rows), default=0)
    for row in rows:
        yield row + [""] * (width - len(row))


def _set_docx_font(run: Any, font_name: str, size_pt: float | None = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size_pt is not None:
        run.font.size = Pt(size_pt)


def _style_docx(document: Any) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def export_docx(blocks: list[ReportBlock], output_path: str | Path) -> Path:
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.start_type = WD_SECTION_START.NEW_PAGE
    _style_docx(document)

    for block in blocks:
        if block.kind == "title":
            paragraph = document.add_heading(block.text, 0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif block.kind == "heading":
            document.add_heading(block.text, level=block.level)
        elif block.kind == "bullet":
            document.add_paragraph(block.text, style="List Bullet")
        elif block.kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block.text)
            _set_docx_font(run, "Consolas", 9)
        elif block.kind == "table" and block.rows:
            normalized_rows = list(_iter_table_rows(block.rows))
            table = document.add_table(rows=len(normalized_rows), cols=len(normalized_rows[0]))
            table.style = "Table Grid"
            table.autofit = True
            for row_index, row in enumerate(normalized_rows):
                for col_index, cell in enumerate(row):
                    target_cell = table.cell(row_index, col_index)
                    target_cell.text = cell
                    for paragraph in target_cell.paragraphs:
                        for run in paragraph.runs:
                            _set_docx_font(run, "Microsoft YaHei", 9.5)
                    if row_index == 0:
                        for paragraph in target_cell.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True
                                run.font.size = Pt(9.5)
            document.add_paragraph()
        else:
            document.add_paragraph(block.text)
    document.save(path)
    return path


def export_report_documents(
    reports_dir: str | Path,
    output_dir: str | Path,
    *,
    basename: str = "bucad_report_summary",
    formats: Iterable[str] = ("docx",),
) -> dict[str, Path]:
    selected = {item.lower() for item in formats}
    if not selected:
        raise ValueError("At least one report format must be selected.")
    unsupported = selected - SUPPORTED_FORMATS
    if unsupported:
        raise ValueError(f"Unsupported report format(s): {', '.join(sorted(unsupported))}")

    blocks = collect_report_blocks(reports_dir)
    output_root = Path(output_dir)
    outputs: dict[str, Path] = {}
    if "docx" in selected:
        outputs["docx"] = export_docx(blocks, output_root / f"{basename}.docx")
    return outputs
